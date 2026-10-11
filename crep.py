"""
Lógica del generador de códigos CREP.

Es un port fiel del Google Apps Script original (no se cambió ninguna regla):
  - A <- W Código Matrícula
  - B <- F Alumno, corregido para BCP (MAYÚS, sin tildes, SIN COMAS, sin símbolos)
  - C <- Q Nro Cuota (texto tal cual, ej. "2-1")
  - D/E/F <- día/mes/año de K (Fecha Venc. Actual) con ajuste de duplicados
  - G <- L Monto Cuota Actual
  - H <- fórmula del código CREP

Filtros previos (antes de pasar al "flujo"):
  - R Moneda: soles  /  dólares (acepta dolares, Dolares, Dólares, etc.)
  - U Saldo Pendiente Dolar distinto de 0 exacto
  - K Fecha Vencimiento Actual >= 01/01/2023
"""
import calendar
import io
import re
import unicodedata
from collections import Counter
from datetime import date
from decimal import Decimal, ROUND_HALF_UP

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- constantes
# Posición (0 = columna A) de cada columna del reporte de INTEGRA
I_ALUMNO = 5      # F
I_FECVENC = 10    # K
I_MONTO = 11      # L  (Monto Cuota Actual)
I_NROCUOTA = 16   # Q
I_MONEDA = 17     # R
I_SALDO = 20      # U  (Saldo Pendiente Dolar)
I_CODMAT = 22     # W

FECHA_MIN = pd.Timestamp("2023-01-01")

# Convenio y fecha de proceso por moneda (tal cual el script; editables aquí)
CONV_SOLES, FPROC_SOLES = "DD21501863341", "20230213"
CONV_DOLARES, FPROC_DOLARES = "DD21511870934", "20220707"

HOJA_FLUJO_S, HOJA_FLUJO_D = "Flujo soles", "Flujo dólares"
HOJA_COD_S, HOJA_COD_D = "Código libre soles", "Código libre dólares"
ENCABEZADOS_CODIGO = ["Código matrícula", "Alumno", "Cuota", "Día", "Mes", "Año", "Monto", "CREP"]


# ---------------------------------------------------------------- utilidades
def quitar_tildes(x) -> str:
    s = unicodedata.normalize("NFD", str(x if x is not None else ""))
    return re.sub(r"[\u0300-\u036f]", "", s)


def norm(x) -> str:
    """MAYÚSCULAS, sin tildes y sin espacios en los bordes (para comparar monedas)."""
    return quitar_tildes(x).upper().strip()


def fix_nombre(x) -> str:
    """Corrige nombre para BCP: MAYÚS, sin tildes, SIN COMAS, sin símbolos."""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    s = quitar_tildes(x).upper()
    s = s.replace(",", " ")
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def num(x) -> float:
    if x is None or x == "" or (isinstance(x, float) and pd.isna(x)):
        return 0.0
    try:
        return float(str(x).replace(",", ""))
    except ValueError:
        return 0.0


def a_texto(x) -> str:
    """str(x).strip(), cuidando que 123.0 no quede como '123.0'."""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    if isinstance(x, float) and x.is_integer():
        return str(int(x))
    return str(x).strip()


# ---------------------------------------------------------------- lectura (por filas)
def iterar_filas(contenido: bytes):
    """
    Devuelve las filas de la primera hoja una por una (sin cargar todo en memoria).
    Usa calamine (rápido); si no está disponible, openpyxl en modo lectura.
    """
    try:
        from python_calamine import CalamineWorkbook
        wb = CalamineWorkbook.from_filelike(io.BytesIO(contenido))
        yield from wb.get_sheet_by_index(0).iter_rows()
    except ImportError:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        yield from wb.worksheets[0].iter_rows(values_only=True)


def a_fecha(v):
    """Convierte el valor de una celda a Timestamp (o None si no se puede)."""
    if v is None or v == "":
        return None
    t = pd.to_datetime(v, errors="coerce")
    return None if pd.isna(t) else t


def saldo_distinto_de_cero(v) -> bool:
    """True si el saldo NO es 0 exacto (vacío o no numérico se conserva, como en el filtro de Sheets)."""
    try:
        return float(v) != 0.0
    except (TypeError, ValueError):
        return True


def validar_columnas(encabezados) -> list:
    """Avisos si el reporte no tiene la estructura esperada (A..W como mínimo)."""
    if len(encabezados) < 23:
        return [f"El archivo tiene {len(encabezados)} columnas; se esperaban al menos 23 (A..W)."]
    esperado = {
        I_ALUMNO: "alumno", I_FECVENC: "vencimiento", I_MONTO: "monto cuota",
        I_NROCUOTA: "cuota", I_MONEDA: "moneda", I_SALDO: "saldo pendiente", I_CODMAT: "matricula",
    }
    avisos = []
    for i, clave in esperado.items():
        nombre = norm(encabezados[i]).lower()
        if clave not in nombre:
            avisos.append(
                f"La columna {get_column_letter(i + 1)} se llama '{encabezados[i]}' "
                f"y se esperaba algo con '{clave}'. Revisa que sea el reporte correcto."
            )
    return avisos


def leer_y_filtrar(contenido: bytes):
    """
    Recorre el reporte UNA vez y se queda solo con las filas que pasan los filtros.
      SOLES   : Moneda = soles   | U (Saldo Pendiente Dolar) != 0 | K (Fecha Venc.) >= 01/01/2023
      DÓLARES : Moneda = dólares | (mismos filtros)
    Mantiene el orden original. Devuelve (flujo_soles, flujo_dolares, monedas, total, avisos).
    """
    kept = {"SOLES": [], "DOLARES": []}
    monedas = Counter()
    encabezados, total = None, 0

    for k, row in enumerate(iterar_filas(contenido)):
        if k == 0:
            encabezados = [str(c) for c in row]
            avisos = validar_columnas(encabezados)
            if len(encabezados) < 23:
                raise ValueError(avisos[0])
            continue
        if all(c is None or c == "" for c in row):
            continue                                   # fila vacía
        row = list(row) + [None] * (len(encabezados) - len(row))
        total += 1
        mon = norm(row[I_MONEDA])
        monedas[mon] += 1
        if mon not in kept:
            continue
        if not saldo_distinto_de_cero(row[I_SALDO]):
            continue
        fec = a_fecha(row[I_FECVENC])
        if fec is None or fec < FECHA_MIN:
            continue
        kept[mon].append(row[:len(encabezados)])

    flujo_s = pd.DataFrame(kept["SOLES"], columns=encabezados)
    flujo_d = pd.DataFrame(kept["DOLARES"], columns=encabezados)
    tabla = (pd.DataFrame(monedas.most_common(), columns=["Moneda", "Filas en el reporte"]))
    return flujo_s, flujo_d, tabla, total, avisos


# ---------------------------------------------------------------- ajuste de fechas
def ajustar_fechas(docs, fechas):
    """
    Port fiel de ajustarFechas(). NO cambia el orden de las filas: devuelve una
    lista paralela donde el elemento i corresponde a la fila i.
    Fechas únicas = intocables. Duplicados (mismo doc + misma fecha) buscan día
    libre hacia adelante; al topar fin de mes retroceden. Nunca cruza de mes.
    """
    n = len(docs)
    adj = list(fechas)

    por_doc = {}
    for i in range(n):
        if fechas[i] is None:
            continue
        por_doc.setdefault(docs[i], []).append(i)

    for pos in por_doc.values():
        por_mes = {}
        for p in pos:
            por_mes.setdefault((fechas[p].year, fechas[p].month), []).append(p)

        for (y, m), ps in por_mes.items():
            last_day = calendar.monthrange(y, m)[1]

            cnt = {}
            for p in ps:
                d = fechas[p].day
                cnt[d] = cnt.get(d, 0) + 1

            occ = {d: True for d, c in cnt.items() if c == 1}
            ordenados = sorted(ps, key=lambda p: fechas[p].day)  # estable, como en JS
            seen = {}

            for p in ordenados:
                day = fechas[p].day
                if cnt[day] > 1 and not seen.get(day):
                    seen[day] = True
                    occ[day] = True
                    asign = day
                elif cnt[day] == 1:
                    asign = day
                else:
                    d2 = day
                    while occ.get(d2) and d2 < last_day:
                        d2 += 1
                    if occ.get(d2) and d2 >= last_day:
                        d2 = last_day
                        while occ.get(d2) and d2 > 1:
                            d2 -= 1
                    occ[d2] = True
                    asign = d2
                adj[p] = date(y, m, asign)
    return adj


# ---------------------------------------------------------------- código CREP
def formula_codigo(r: int, conv: str, fproc: str) -> str:
    """Fórmula del código para la columna H (idéntica a la del script)."""
    doc = f'RIGHT(REPT(" ",14)&A{r},14)'
    nom = f'LEFT(B{r}&REPT(" ",40),40)'
    M = f'IF(LEN(C{r})=4,LEFT(C{r},2),"0"&LEFT(C{r},1))'
    cuota = f'"1"&{M}&"0"&RIGHT(C{r},1)&"XXXXXX"'
    fven = f'F{r}&RIGHT("0"&E{r},2)&RIGHT("0"&D{r},2)'
    ent = f'RIGHT(REPT("0",13)&INT(G{r}),13)'
    cent = f'TEXT(ROUND(MOD(G{r},1)*100,0),"00")'
    monto = f'{ent}&{cent}&REPT("0",24)&"A"'
    return (f'="{conv}"&{doc}&{nom}&{cuota}&"                   "&"{fproc}"&{fven}&{monto}')


def crep_valor(doc, nombre, cuota, d, e, f, g, conv, fproc) -> str:
    """Mismo código CREP calculado en Python (vista previa, verificación y opción 'H como valor')."""
    doc_s = ((" " * 14) + doc)[-14:]
    nom_s = (nombre + " " * 40)[:40]
    M = cuota[:2] if len(cuota) == 4 else "0" + cuota[:1]
    cuota_s = "1" + M + "0" + cuota[-1:] + "XXXXXX"
    fven = f"{f}" + ("0" + f"{e}")[-2:] + ("0" + f"{d}")[-2:]
    centavos = int((Decimal(str(g)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    ent = (("0" * 13) + str(centavos // 100))[-13:]
    cent = f"{centavos % 100:02d}"
    return f"{conv}{doc_s}{nom_s}{cuota_s}{' ' * 19}{fproc}{fven}{ent}{cent}{'0' * 24}A"


# ---------------------------------------------------------------- transformación
def generar_filas(flujo: pd.DataFrame, es_soles: bool):
    """
    Cada fila del flujo -> su fila de códigos (mismo orden).
    Devuelve (filas, avisos). Cada fila: [A, B, C, D, E, F, G].
    """
    avisos = []
    n = len(flujo)
    docs = [a_texto(v) for v in flujo.iloc[:, I_CODMAT]]
    fechas = []
    for v in pd.to_datetime(flujo.iloc[:, I_FECVENC], errors="coerce"):
        fechas.append(None if pd.isna(v) else v.date())

    salida = []
    cuotas_raras = 0
    for i in range(n):
        cuota = a_texto(flujo.iloc[i, I_NROCUOTA])
        if not re.fullmatch(r"\d{1,2}-\d", cuota):
            cuotas_raras += 1
        salida.append([
            docs[i],
            fix_nombre(flujo.iloc[i, I_ALUMNO]),
            cuota,
            "", "", "",
            num(flujo.iloc[i, I_MONTO]),
        ])
    if cuotas_raras:
        avisos.append(
            f"{cuotas_raras} filas tienen un 'Nro Cuota' con formato distinto a '2-1' "
            "(la fórmula del CREP asume ese formato). Revísalas."
        )

    adj = ajustar_fechas(docs, fechas)
    for i, f in enumerate(adj):
        if f:
            salida[i][3], salida[i][4], salida[i][5] = f.day, f.month, f.year
    return salida, avisos


# ---------------------------------------------------------------- Excel de salida
def _escribir_flujo(wb, nombre, flujo):
    ws = wb.create_sheet(nombre)
    ws.append(list(flujo.columns))
    for c in ws[1]:
        c.font = Font(bold=True)
    datos = flujo.astype(object).where(flujo.notna(), None)
    for fila in datos.values.tolist():
        ws.append(fila)
    for j, col in enumerate(flujo.columns, start=1):
        if pd.api.types.is_datetime64_any_dtype(flujo[col]):
            for r in range(2, len(flujo) + 2):
                ws.cell(row=r, column=j).number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A2"


def _escribir_codigos(wb, nombre, filas, conv, fproc, h_como_valor):
    ws = wb.create_sheet(nombre)
    ws.append(ENCABEZADOS_CODIGO)
    for c in ws[1]:
        c.font = Font(bold=True)
    for k, fila in enumerate(filas):
        r = k + 2
        if h_como_valor:
            h = crep_valor(fila[0], fila[1], fila[2], fila[3], fila[4], fila[5], fila[6], conv, fproc)
        else:
            h = formula_codigo(r, conv, fproc)
        ws.append(fila + [h])
        ws.cell(row=r, column=1).number_format = "@"   # A texto
        ws.cell(row=r, column=3).number_format = "@"   # C texto ("2-1")
    ws.freeze_panes = "A2"
    for col, ancho in zip("ABCDEFGH", (18, 42, 8, 6, 6, 8, 12, 140)):
        ws.column_dimensions[col].width = ancho


def construir_excel(flujo_s, filas_s, flujo_d, filas_d, h_como_valor=False) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _escribir_flujo(wb, HOJA_FLUJO_S, flujo_s)
    _escribir_flujo(wb, HOJA_FLUJO_D, flujo_d)
    _escribir_codigos(wb, HOJA_COD_S, filas_s, CONV_SOLES, FPROC_SOLES, h_como_valor)
    _escribir_codigos(wb, HOJA_COD_D, filas_d, CONV_DOLARES, FPROC_DOLARES, h_como_valor)
    wb.calculation.fullCalcOnLoad = True   # Excel recalcula las fórmulas al abrir
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------- proceso completo
def procesar(contenido: bytes, h_como_valor: bool = False) -> dict:
    flujo_s, flujo_d, monedas, total, avisos = leer_y_filtrar(contenido)
    filas_s, av_s = generar_filas(flujo_s, es_soles=True)
    filas_d, av_d = generar_filas(flujo_d, es_soles=False)
    avisos += [f"Soles: {a}" for a in av_s] + [f"Dólares: {a}" for a in av_d]
    if len(flujo_s) == 0:
        avisos.append("No quedó ninguna fila de SOLES después de filtrar.")
    if len(flujo_d) == 0:
        avisos.append("No quedó ninguna fila de DÓLARES después de filtrar.")

    excel = construir_excel(flujo_s, filas_s, flujo_d, filas_d, h_como_valor)

    def vista_previa(filas, conv, fproc, k=30):
        t = pd.DataFrame(filas[:k], columns=ENCABEZADOS_CODIGO[:7])
        t["CREP"] = [crep_valor(*f, conv, fproc) for f in filas[:k]]
        return t

    return {
        "excel": excel,
        "n_reporte": total,
        "n_soles": len(flujo_s),
        "n_dolares": len(flujo_d),
        "monedas": monedas,
        "avisos": avisos,
        "vista_soles": vista_previa(filas_s, CONV_SOLES, FPROC_SOLES),
        "vista_dolares": vista_previa(filas_d, CONV_DOLARES, FPROC_DOLARES),
    }
